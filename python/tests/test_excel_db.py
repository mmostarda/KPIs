import datetime as dt
import hashlib
import os
import shutil
import subprocess
import zipfile

import pandas as pd
import pytest
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.packaging.custom import StringProperty

from kpimeta import excel_db
from kpimeta.batch import batch_from_frame, compose
from kpimeta.excel_db import (DBChangedError, DBError, DBLockedError, add_support_value, overwrite_metadata,
                              read_db, write_batch)
from kpimeta.issues import ERROR, WARNING
from kpimeta.lola import read_lola_files
from kpimeta.mapping import load_mapping
from kpimeta.validate import KpiDecision


def _digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _batch(schema, settings, rows=2, kpi="KPI_A"):
    frame = pd.DataFrame([{"IDName": f"new {i}", "Date": f"1{i}/04/2026", "Brand": "Brand X",
                           "StartSpeed": 80 + i, kpi: 0.5 + i} for i in range(rows)], dtype=object)
    return batch_from_frame(frame, schema, settings, "excel")


def _edit(path, fn):
    wb = load_workbook(path)
    fn(wb)
    wb.save(path)


def test_demo_db_is_clean(db_path, schema, settings):
    snapshot = read_db(db_path, schema, settings)
    assert snapshot.issues == []
    assert snapshot.next_id == 3 and snapshot.kpi_names == ["KPI_A", "KPI_B", "KPI_C"]
    assert snapshot.options(schema)["B_F_Type"] == ["Type 1", "Type 2"]
    assert list(snapshot.templates["IDName_Template"]) == ["PRJ1_VEH1", "PRJ2_VEH2"]
    assert snapshot.metadata_row(2)["IDName"] == "Demo 2"


def test_full_lola_batch_write(demo, schema, settings):
    _edit(demo["db"], lambda wb: setattr(wb["KPIs"]["D3"], "number_format", "0.000"))
    snapshot = read_db(demo["db"], schema, settings)
    frame, _ = read_lola_files(demo["lola"], load_mapping(demo["mapping"]), settings)
    batch = compose(batch_from_frame(frame, schema, settings, "lola", snapshot.metadata.headers), schema,
                    settings, {"IDName": "LoLa batch"}, templates=snapshot.templates,
                    infer_template=True, infer_filename=True)
    batch.conv_issues.clear()  # the user emptied the 'n.a.' cell
    result = write_batch(demo["db"], batch, schema, settings, {"KPI_D": KpiDecision("add")})
    assert result.ids == list(range(3, 11)) and result.new_kpi_columns == ["KPI_D"]
    assert result.backup.exists() and result.backup.parent.parent == settings.backup_dir

    after = read_db(demo["db"], schema, settings)
    assert after.issues == [] and after.next_id == 11
    kpis = after.kpis.df
    assert list(kpis["ID"]) == list(range(1, 11))
    assert kpis.iloc[2]["KPI_A"] == 0.123456789 and kpis.iloc[2]["KPI_D"] == 1
    assert kpis.iloc[2]["Date"] == "15/03/2026"  # text dd/mm/yyyy (date_storage = "text")
    meta = after.metadata.df
    assert meta.iloc[2]["B_F_Cal_Date"] == "10/01/2026" and meta.iloc[7]["Brand"] == "Brand Y"
    wb = load_workbook(demo["db"])
    assert wb["KPIs"].tables["TableKPI"].ref == "A1:G11"
    assert wb["KPIs"]["D11"].number_format == "0.000"  # new rows take the style of the last row
    assert not [p for p in demo["db"].parent.iterdir() if p.name.startswith(".kpimeta-tmp")]


def test_refused_when_open_in_excel(db_path, schema, settings):
    (db_path.parent / f"~${db_path.name}").write_text("lock")
    before = _digest(db_path)
    assert read_db(db_path, schema, settings).lock
    with pytest.raises(DBLockedError):
        write_batch(db_path, _batch(schema, settings), schema, settings)
    assert _digest(db_path) == before


def test_aborted_if_db_changes_meanwhile(db_path, schema, settings, monkeypatch):
    original_verify = excel_db._verify

    def verify_then_touch(*args):
        messages = original_verify(*args)
        os.utime(db_path, ns=(db_path.stat().st_atime_ns, db_path.stat().st_mtime_ns + 5_000_000_000))
        return messages

    monkeypatch.setattr(excel_db, "_verify", verify_then_touch)
    before = _digest(db_path)
    with pytest.raises(DBChangedError):
        write_batch(db_path, _batch(schema, settings), schema, settings)
    assert _digest(db_path) == before


def test_failure_leaves_original_untouched(db_path, schema, settings, monkeypatch):
    def boom(*args, **kwargs):
        raise RuntimeError("simulated crash while writing")

    monkeypatch.setattr(excel_db, "_append_rows", boom)
    before = _digest(db_path)
    with pytest.raises(RuntimeError):
        write_batch(db_path, _batch(schema, settings), schema, settings)
    assert _digest(db_path) == before
    assert not [p for p in db_path.parent.iterdir() if p.name.startswith(".kpimeta-tmp")]


def test_damaged_or_non_excel_file(tmp_path, schema, settings):
    damaged = tmp_path / "damaged.xlsx"
    damaged.write_bytes(b"not a zip")
    with pytest.raises(DBError, match="damaged.xlsx: non è un file Excel valido"):
        read_db(damaged, schema, settings)
    other_zip = tmp_path / "other.xlsx"
    with zipfile.ZipFile(other_zip, "w") as archive:
        archive.writestr("readme.txt", "a zip that is not a workbook")
    with pytest.raises(DBError, match="other.xlsx: non è un file Excel valido"):
        read_db(other_zip, schema, settings)
    with pytest.raises(DBError, match="other.xlsx: non è un file Excel valido"):
        write_batch(other_zip, _batch(schema, settings), schema, settings)


def test_inconsistent_db_blocks_writing(db_path, schema, settings):
    def drop_kpi_row(wb):
        ws = wb["KPIs"]
        ws.delete_rows(3)
        ws.tables["TableKPI"].ref = "A1:F2"
    _edit(db_path, drop_kpi_row)
    snapshot = read_db(db_path, schema, settings)
    assert any(i.severity == ERROR and "senza riga KPI" in i.message for i in snapshot.issues)
    with pytest.raises(DBError, match="senza riga KPI"):
        write_batch(db_path, _batch(schema, settings), schema, settings)


def test_formula_inside_table_blocks(db_path, schema, settings):
    _edit(db_path, lambda wb: wb["KPIs"].__setitem__("F2", "=D2*2"))
    snapshot = read_db(db_path, schema, settings)
    assert any(i.severity == ERROR and "formula" in i.message for i in snapshot.issues)
    with pytest.raises(DBError, match="formula"):
        write_batch(db_path, _batch(schema, settings), schema, settings)


def test_formula_elsewhere_is_a_warning(db_path, schema, settings):
    _edit(db_path, lambda wb: wb["Brand"].__setitem__("E1", "=1+1"))
    snapshot = read_db(db_path, schema, settings)
    assert snapshot.writable and any(i.severity == WARNING and "formule" in i.message for i in snapshot.issues)
    write_batch(db_path, _batch(schema, settings), schema, settings)


def test_data_below_the_table_blocks(db_path, schema, settings):
    _edit(db_path, lambda wb: wb["KPIs"].__setitem__("A5", "note"))
    with pytest.raises(DBError, match="non è vuota"):
        write_batch(db_path, _batch(schema, settings, rows=3), schema, settings)


def test_totals_row_blocks(db_path, schema, settings):
    def totals(wb):
        table = wb["KPIs"].tables["TableKPI"]
        table.ref = "A1:F4"
        table.totalsRowCount = 1
        wb["KPIs"]["A4"] = "Totale"
    _edit(db_path, totals)
    assert not read_db(db_path, schema, settings).writable


def test_chart_and_custom_properties_survive(db_path, schema, settings):
    def add(wb):
        ws = wb["Brand"]
        chart = BarChart()
        chart.add_data(Reference(ws, min_col=1, min_row=1, max_row=3))
        ws.add_chart(chart, "E2")
        wb.custom_doc_props.append(StringProperty(name="MSIP_Label_demo_Enabled", value="true"))
    _edit(db_path, add)
    snapshot = read_db(db_path, schema, settings)
    assert snapshot.writable and any("grafici" in i.message for i in snapshot.issues)
    write_batch(db_path, _batch(schema, settings), schema, settings)
    names = zipfile.ZipFile(db_path).namelist()
    assert any(n.startswith("xl/charts/") for n in names)
    assert [p.name for p in load_workbook(db_path).custom_doc_props] == ["MSIP_Label_demo_Enabled"]


def test_excel_date_storage(db_path, schema, settings):
    settings.date_storage = "excel"
    write_batch(db_path, _batch(schema, settings, rows=1), schema, settings)
    cell = load_workbook(db_path)["Metadata"]["C4"]
    assert cell.value == dt.datetime(2026, 4, 10) and cell.number_format == "dd/mm/yyyy"


def test_first_write_into_empty_tables(db_path, schema, settings):
    def empty(wb):
        for sheet, table, width in (("Metadata", "TableMet", len(schema.columns)), ("KPIs", "TableKPI", 6)):
            ws = wb[sheet]
            ws.delete_rows(2, 2)
            ws.tables[table].ref = f"A1:{ws.cell(1, width).column_letter}2"  # one blank data row
    _edit(db_path, empty)
    snapshot = read_db(db_path, schema, settings)
    assert snapshot.next_id == 1 and snapshot.writable
    result = write_batch(db_path, _batch(schema, settings, rows=2), schema, settings)
    assert result.ids == [1, 2]
    assert load_workbook(db_path)["KPIs"].tables["TableKPI"].ref == "A1:F3"


def test_overwrite_updates_metadata_and_kpi_row(db_path, schema, settings):
    result = overwrite_metadata(db_path, 2, {"IDName": "Renamed", "Date": dt.date(2026, 5, 1), "Brand": "Brand Y"},
                                schema, settings)
    assert "IDName" in result.messages[0] and "IDName" in result.messages[1]
    after = read_db(db_path, schema, settings)
    assert after.metadata_row(2)["IDName"] == "Renamed" and after.metadata_row(1)["IDName"] == "Demo 1"
    assert list(after.kpis.df["IDName"]) == ["Demo 1", "Renamed"]
    assert after.kpis.df.iloc[1]["Date"] == "01/05/2026"
    with pytest.raises(DBError, match="ID 99"):
        overwrite_metadata(db_path, 99, {"IDName": "x"}, schema, settings)


def test_add_support_value(db_path, schema, settings):
    add_support_value(db_path, "Brand", {"Brand": "Brand Z"}, settings)
    snapshot = read_db(db_path, schema, settings)
    assert snapshot.options(schema)["Brand"] == ["Brand X", "Brand Y", "Brand Z"]
    ws = load_workbook(db_path)["Brand"]
    assert ws["A4"].value == 3 and ws.tables["T_Brand"].ref == "A1:B4"


def test_matlab_datenum_dates_are_reported(db_path, schema, settings):
    _edit(db_path, lambda wb: wb["Metadata"].__setitem__("AH2", 739983))  # B_F_Cal_Date
    snapshot = read_db(db_path, schema, settings)
    assert any("datenum" in i.message for i in snapshot.issues)


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice non installato")
def test_libreoffice_reads_the_written_file(db_path, schema, settings, tmp_path):
    write_batch(db_path, _batch(schema, settings), schema, settings)
    out = tmp_path / "lo"
    subprocess.run(["soffice", f"-env:UserInstallation=file://{tmp_path}/profile", "--headless", "--convert-to",
                    "csv:Text - txt - csv (StarCalc):59,34,UTF8,1,,0,false,true,false,false,false,-1",
                    "--outdir", str(out), str(db_path)], capture_output=True, timeout=180, check=True)
    lines = (out / f"{db_path.stem}-KPIs.csv").read_text(encoding="utf-8").splitlines()
    assert lines[0] == "ID;IDName;Date;KPI_A;KPI_B;KPI_C" and lines[-1].startswith("4;new 1;11/04/2026;1.5")
