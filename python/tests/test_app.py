"""Streamlit flow driven with AppTest (the editable grid itself is not scriptable here)."""

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from streamlit.testing.v1 import AppTest

APP = str(Path(__file__).resolve().parent.parent / "app.py")


def _button(at, label):
    matches = [b for b in at.button if b.label == label]
    assert matches, f"pulsante '{label}' non trovato: {[b.label for b in at.button]}"
    return matches[0]


def _no_exception(at):
    assert not at.exception, [e.value for e in at.exception]


@pytest.fixture()
def app(tmp_path, monkeypatch, demo):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "appdata"))
    monkeypatch.delenv("DISPLAY", raising=False)
    at = AppTest.from_file(APP, default_timeout=120)
    at.run()
    _no_exception(at)
    return at


def test_lola_flow_writes_the_db(app, demo):
    at = app
    at.text_input(key="db_path").set_value(str(demo["db"]))
    _button(at, "Connetti / rileggi").click().run()
    _no_exception(at)
    assert any("DB coerente" in s.value for s in at.success)

    at.button(key="next::1 · DB").click().run()
    at.text_input(key="mapping_path").set_value(str(demo["mapping"])).run()
    at.text_input(key="lola_dir").set_value(str(Path(demo["lola"][0]).parent)).run()
    _button(at, "Leggi i file LoLa").click().run()
    _no_exception(at)
    assert any("8 righe lette" in s.value for s in at.success)

    at.button(key="next::2 · Sorgente").click().run()
    at.selectbox(key="template_choice").set_value("PRJ1_VEH1").run()
    _button(at, "Applica template").click().run()
    assert at.selectbox(key="f::Brand").value == "Brand X"
    at.text_input(key="f::IDName").set_value("Da AppTest").run()
    at.checkbox(key="infer_template").check().run()
    at.button(key="next::3 · Metadati").click().run()
    _no_exception(at)

    errors = [m for m in at.metric if m.label == "Errori"][0]
    assert errors.value == "1"  # the 'n.a.' KPI value
    _button(at, "Accetta come vuoti i valori non validi").click().run()
    assert [m for m in at.metric if m.label == "Errori"][0].value == "0"

    at.button(key="next::4 · Revisione").click().run()
    _no_exception(at)
    assert at.selectbox(key="kpi::KPI_D").value == "— scegli —"
    _button(at, "Tutti: nuova colonna").click().run()
    assert any("ID 3–10" in i.value for i in at.info)
    _button(at, "Scrivi nel DB").click().run()
    _no_exception(at)
    assert any("Scritti gli ID 3–10" in s.value for s in at.success)

    wb = load_workbook(demo["db"])
    assert wb["KPIs"].tables["TableKPI"].ref == "A1:G11"
    assert wb["Metadata"]["B4"].value == "Da AppTest"


def test_failed_read_drops_the_previous_rows(app, demo):
    at = app
    at.text_input(key="db_path").set_value(str(demo["db"]))
    _button(at, "Connetti / rileggi").click().run()
    at.button(key="next::1 · DB").click().run()
    folder = Path(demo["lola"][0]).parent
    at.text_input(key="mapping_path").set_value(str(demo["mapping"])).run()
    at.text_input(key="lola_dir").set_value(str(folder)).run()
    _button(at, "Leggi i file LoLa").click().run()
    assert not at.button(key="next::2 · Sorgente").disabled

    (folder / "broken.xlsx").write_bytes(b"not an Excel file")
    at.run()
    at.multiselect(key=f"lola_files::{folder}").select("broken.xlsx").run()
    _button(at, "Leggi i file LoLa").click().run()
    _no_exception(at)
    assert any("broken.xlsx: non è un file Excel valido" in e.value for e in at.error)
    assert at.button(key="next::2 · Sorgente").disabled  # the rows of the first read are gone


def test_form_survives_step_changes(app, demo):
    at = app
    at.text_input(key="db_path").set_value(str(demo["db"]))
    _button(at, "Connetti / rileggi").click().run()
    at.radio(key="step").set_value("2 · Sorgente").run()
    at.radio(key="source").set_value("manual").run()
    at.button(key="next::2 · Sorgente").click().run()
    at.text_input(key="f::IDName").set_value("persistente").run()
    at.radio(key="step").set_value("1 · DB").run()
    at.radio(key="step").set_value("3 · Metadati").run()
    assert at.text_input(key="f::IDName").value == "persistente"


def test_edit_flow(app, demo):
    at = app
    at.text_input(key="db_path").set_value(str(demo["db"]))
    _button(at, "Connetti / rileggi").click().run()
    at.radio(key="step").set_value("2 · Sorgente").run()
    at.radio(key="source").set_value("edit").run()
    at.number_input(key="edit_id_input").set_value(2).run()
    _button(at, "Carica nel form").click().run()
    at.button(key="next::2 · Sorgente").click().run()
    assert at.text_input(key="f::IDName").value == "Demo 2"
    at.text_input(key="f::IDName").set_value("Demo 2 corretto").run()
    at.button(key="next::3 · Metadati").click().run()
    at.button(key="next::4 · Revisione").click().run()
    _button(at, "Sovrascrivi l'ID 2").click().run()
    _no_exception(at)
    ws = load_workbook(demo["db"])["KPIs"]
    assert ws["B3"].value == "Demo 2 corretto"  # KPI row kept in sync


def test_review_problem_columns_and_support_values(app, demo, tmp_path):
    at = app
    source = tmp_path / "righe.xlsx"  # a Brand missing from the Brand sheet of the DB + an invalid KPI
    wb = Workbook()
    wb.active.append(["IDName", "Brand", "FileName", "KPI_A"])
    wb.active.append(["Da Excel", "Brand Z", "A_B_20260101_x_Sport_BrakeTest_DriverA_1.mf4", 1.5])
    wb.active.append(["Da Excel", "Brand Z", "A_B_20260102_x_Sport_BrakeTest_DriverA_2.mf4", "n.a."])
    wb.save(source)
    at.text_input(key="db_path").set_value(str(demo["db"]))
    _button(at, "Connetti / rileggi").click().run()
    at.radio(key="step").set_value("2 · Sorgente").run()
    at.radio(key="source").set_value("excel").run()
    at.text_input(key="excel_path").set_value(str(source)).run()
    _button(at, "Leggi il file").click().run()
    at.button(key="next::2 · Sorgente").click().run()
    at.button(key="next::3 · Metadati").click().run()
    _no_exception(at)

    _button(at, "Mostra solo colonne con problemi").click().run()
    assert set(at.multiselect(key="review_cols").value) == {"Brand", "KPI_A"}
    _button(at, "Mostra le colonne predefinite").click().run()
    assert "IDName" in at.multiselect(key="review_cols").value

    # the missing value is added to the support sheet from the review, without going back
    assert at.text_input(key="add::review::Brand Z::Brand::Brand").value == "Brand Z"
    at.button(key="addbtn::review::Brand Z::Brand").click().run()
    _no_exception(at)
    assert any("'Brand Z' aggiunto al foglio Brand" in s.value for s in at.success)
    assert not [b for b in at.button if b.key == "addbtn::review::Brand Z::Brand"]  # no longer missing
    brands = [r[1] for r in load_workbook(demo["db"])["Brand"].iter_rows(min_row=2, values_only=True)]
    assert brands == ["Brand X", "Brand Y", "Brand Z"]
