import datetime as dt

import pytest
from openpyxl import Workbook

from kpimeta.lola import LolaError, export_all_db, read_lola_file, read_lola_files
from kpimeta.mapping import load_mapping


def test_demo_outputs(demo, settings):
    mapping = load_mapping(demo["mapping"])
    frame, reports = read_lola_files(demo["lola"], mapping, settings)
    assert [r.data_rows for r in reports] == [5, 3]
    first = frame[frame["_source_file"] == "LoLa_output_1.xlsx"]
    # merged "File Name" filled down, statistics row ("Mean") and separators ignored
    assert list(first["Repetition_ID"]) == [1, 2, 3, 1, 2]
    assert first["FileName"].nunique() == 2
    # numbers keep full precision (MATLAB num2str kept ~5 significant digits)
    assert first.iloc[0]["KPI_A"] == 0.123456789
    # text left untouched here: the batch validation reports it
    assert first.iloc[0]["KPI_C"] == "n.a."


def _write(path, rows):
    wb = Workbook()
    for row in rows:
        wb.active.append(row)
    wb.save(path)
    return path


def test_layout_rules(tmp_path, settings):
    mapping_path = tmp_path / "m.csv"
    mapping_path.write_text("KPI_LoLa_Name;KPI_DB_Name;Type\nFile Name;FileName;text\n"
                            "UpdateRate;UpdateRate;\nTest Date;TestDate;date\nMissing col;Nope;\n",
                            encoding="utf-8")
    path = _write(tmp_path / "lola.xlsx", [
        [None, "File Name", "UpdateRate", "Test Date", "Extra"],
        [1, "A_B_20260101_x_M_N_D_1.mf4", 5.5, "05/03/2026", 1],
        ["2", None, 6.5, dt.datetime(2026, 3, 6), 2],  # numeric text in column A counts (str2double)
        [None, None, None, None, None],  # separator: resets the fill-down
        [1, None, 7.5, None, None],  # no File Name after a separator -> dropped
        ["Std", "x", 1, None, None],  # non numeric column A -> ignored
    ])
    frame, report = read_lola_file(path, load_mapping(mapping_path), settings)
    assert list(frame["Repetition_ID"]) == [1, 2]
    assert list(frame["FileName"]) == ["A_B_20260101_x_M_N_D_1.mf4"] * 2
    # "UpdateRate" contains "date": MATLAB treated it as a date column and emptied it
    assert list(frame["UpdateRate"]) == [5.5, 6.5]
    assert list(frame["TestDate"]) == [dt.date(2026, 3, 5), dt.date(2026, 3, 6)]
    assert report.dropped_without_filename == 1
    assert report.lola_names_missing == ["Missing col"]
    assert report.unmapped_headers == ["Extra"]


def test_export_all_db(demo, settings, tmp_path):
    frame, _ = read_lola_files(demo["lola"], load_mapping(demo["mapping"]), settings)
    out = export_all_db(frame, tmp_path / "ALL_DB.xlsx")
    from openpyxl import load_workbook

    ws = load_workbook(out).active
    header = [c.value for c in ws[1]]
    assert header[0] == "Repetition_ID" and "_source_file" not in header
    assert ws.max_row == len(frame) + 1


def test_damaged_file(demo, settings, tmp_path):
    broken = tmp_path / "broken.xlsx"
    broken.write_bytes(b"not an Excel file")
    with pytest.raises(LolaError, match="broken.xlsx: non è un file Excel valido"):
        read_lola_files([broken], load_mapping(demo["mapping"]), settings)
