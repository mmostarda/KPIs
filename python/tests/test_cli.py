from openpyxl import load_workbook

from kpimeta.cli import main


def test_end_to_end_cli(tmp_path, monkeypatch, capsys):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "appdata"))
    folder = tmp_path / "demo"
    assert main(["demo", str(folder)]) == 0
    db = folder / "KPI_DB_demo.xlsx"
    lola = sorted(str(p) for p in (folder / "lola").glob("*.xlsx"))
    mapping = str(folder / "lola_mapping_demo.csv")
    assert main(["doctor", str(db)]) == 0

    all_db = folder / "ALL_DB.xlsx"
    assert main(["lola", *lola, "-m", mapping, "-o", str(all_db)]) == 0
    assert load_workbook(all_db).active["A1"].value == "Repetition_ID"

    batch = folder / "Metadata_ALL_DB.xlsx"
    assert main(["prepare", "--lola", *lola, "-m", mapping, "--db", str(db), "--template", "PRJ1_VEH1",
                 "--set", "IDName=CLI", "--infer-filename", "--infer-template", "-o", str(batch)]) == 0
    capsys.readouterr()
    assert main(["check", str(batch), "--db", str(db)]) == 1  # 'n.a.' + KPI_D without a decision
    assert "KPI_D" in capsys.readouterr().out

    wb = load_workbook(batch)
    for row in wb.active.iter_rows():
        for cell in row:
            if cell.value == "n.a.":
                cell.value = None
    wb.save(batch)
    assert main(["write", str(batch), "--db", str(db), "--new-kpi", "add", "--dry-run"]) == 0
    assert main(["write", str(batch), "--db", str(db), "--new-kpi", "add", "--yes"]) == 0
    assert "Scritti gli ID 3-10" in capsys.readouterr().out
    assert main(["write", str(batch), "--db", str(db), "--new-kpi", "add", "--yes"]) == 1  # duplicates


def test_convert_mapping(tmp_path, capsys):
    source = tmp_path / "m.csv"
    source.write_text("KPI_LoLa_Name,KPI_DB_Name\nFile Name,FileName\n", encoding="utf-8")
    assert main(["convert-mapping", str(source), "-o", str(tmp_path / "out.csv")]) == 0
    assert (tmp_path / "out.csv").read_text(encoding="utf-8").startswith("KPI_LoLa_Name;KPI_DB_Name;Type")


def test_errors_are_reported_without_traceback(tmp_path, capsys):
    assert main(["doctor", str(tmp_path / "missing.xlsx")]) == 2
    assert "ERRORE" in capsys.readouterr().err


def test_lola_wildcards_skip_intermediate_files(tmp_path, monkeypatch, capsys):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "appdata"))
    folder = tmp_path / "demo"
    main(["demo", str(folder)])
    mapping = str(folder / "lola_mapping_demo.csv")
    pattern = str(folder / "lola" / "*.xlsx")  # expanded by kpimeta (cmd.exe does not do it)
    for _ in range(2):  # the second run must not read the ALL_DB.xlsx written by the first
        capsys.readouterr()
        assert main(["lola", pattern, "-m", mapping]) == 0
        assert "8 righe scritte" in capsys.readouterr().out
    assert main(["lola", str(folder / "none_*.xlsx"), "-m", mapping]) == 2
    assert "nessun file corrisponde" in capsys.readouterr().err
