"""The review grid in a real browser: its edits live in the frontend, which AppTest does not
run (a frontend rule once erased every edit made in the grid). Opt-in, about 30 s:

    KPIMETA_BROWSER_TEST=1 pytest tests/test_browser.py

Needs `pip install playwright` and `playwright install chromium`
(or KPIMETA_CHROMIUM=<path of an existing Chromium>).
"""

from __future__ import annotations

import os
import socket
import subprocess
import sys
import time
import urllib.request
from pathlib import Path

import pytest
from openpyxl import load_workbook

pytestmark = pytest.mark.skipif(not os.environ.get("KPIMETA_BROWSER_TEST"),
                                reason="test nel browser: impostare KPIMETA_BROWSER_TEST=1")

ROOT = Path(__file__).resolve().parent.parent


def _free_port() -> int:
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        return s.getsockname()[1]


@pytest.fixture()
def server(tmp_path):
    port = _free_port()
    env = {k: v for k, v in os.environ.items() if k != "DISPLAY"}  # no "Sfoglia…" buttons
    env["LOCALAPPDATA"] = str(tmp_path / "appdata")
    process = subprocess.Popen([sys.executable, "-m", "streamlit", "run", "app.py", "--server.headless", "true",
                                "--server.port", str(port)], cwd=ROOT, env=env,
                               stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    url = f"http://localhost:{port}"
    no_proxy = urllib.request.build_opener(urllib.request.ProxyHandler({}))
    try:
        deadline = time.monotonic() + 60
        while True:
            try:
                no_proxy.open(f"{url}/_stcore/health", timeout=2)
                break
            except OSError:
                if process.poll() is not None or time.monotonic() > deadline:
                    raise RuntimeError("Streamlit non è partito") from None
                time.sleep(0.5)
        yield url
    finally:
        process.terminate()
        process.wait(timeout=30)


def _idle(page) -> None:
    page.wait_for_timeout(300)
    page.wait_for_selector('[data-testid="stStatusWidget"]', state="detached", timeout=30000)  # "Running…"


def _click(page, name: str) -> None:
    page.get_by_role("button", name=name, exact=True).first.click()
    _idle(page)


def _fill(page, label: str, value) -> None:
    box = page.get_by_label(label, exact=True).first
    box.fill(str(value))
    box.press("Enter")
    _idle(page)


def _metric(page, label: str):
    return page.locator('[data-testid="stMetric"]').filter(has_text=label).locator('[data-testid="stMetricValue"]')


def _chips(page) -> list[str]:
    """Columns chosen in 'Colonne visualizzate', in grid order."""
    return [c.get_attribute("aria-label")
            for c in page.locator('[data-testid="stMultiSelectTagsContainer"] [data-tag]').all()]


def _edit_cell(page, row: int, target: int, text: str, grid: int = 0) -> None:
    """Select the first cell of the row, move right to the column `target` (0 = first column of
    the grid) and type the value (typing on a selected cell opens its editor). Each step waits
    for the grid focus: a key pressed while the grid is still busy can be lost."""
    table = page.locator('[data-testid="stDataFrame"]').nth(grid)
    table.scroll_into_view_if_needed()
    box = table.bounding_box()
    x, y = box["x"] + 20, box["y"] + 35 * row + 17  # header and rows are 35 px high
    page.mouse.click(x, y)
    page.wait_for_timeout(300)
    if page.evaluate("() => document.activeElement?.id") != f"glide-cell-0-{row - 1}":
        page.mouse.click(x, y)  # while another grid has a selection, the first click only activates this one
    for position in range(target + 1):
        if position:
            page.keyboard.press("ArrowRight")
        page.wait_for_function(f"() => document.activeElement?.id === 'glide-cell-{position}-{row - 1}'")
    page.wait_for_timeout(300)  # the grid scrolls to the selected cell
    page.keyboard.type(text[0])
    editor = page.locator(".gdg-input")  # input for numbers, textarea for text
    editor.wait_for()
    editor.fill(text)
    editor.press("Enter")
    _idle(page)
    page.wait_for_timeout(1500)  # the grid may send a second state after the rerun (the bug)
    _idle(page)


def test_grid_edit_reaches_the_db(server, demo):
    sync_api = pytest.importorskip("playwright.sync_api")
    expect = sync_api.expect
    with sync_api.sync_playwright() as p:
        browser = p.chromium.launch(executable_path=os.environ.get("KPIMETA_CHROMIUM") or None)
        page = browser.new_page(viewport={"width": 1500, "height": 950})
        page.set_default_timeout(30000)
        try:
            page.goto(server)
            expect(page.get_by_text("KPI metadata → DB")).to_be_visible(timeout=60000)
            _fill(page, "Percorso del DB", demo["db"])
            _click(page, "Connetti / rileggi")
            expect(page.get_by_text("DB coerente")).to_be_visible()
            _click(page, "Avanti →")
            _fill(page, "Mapping LoLa → DB (CSV)", demo["mapping"])
            _fill(page, "Cartella con gli output LoLa", Path(demo["lola"][0]).parent)
            _click(page, "Leggi i file LoLa")
            expect(page.get_by_text("8 righe lette")).to_be_visible()
            _click(page, "Avanti →")
            page.get_by_label("Template (Metadata_Templates)").first.click()
            page.get_by_role("option", name="PRJ1_VEH1").click()
            _idle(page)
            _click(page, "Applica template")
            _fill(page, "ID Name", "Dal browser")
            _click(page, "Avanti: componi le righe →")

            expect(_metric(page, "Errori")).to_have_text("1")  # 'n.a.' in KPI_C of row 1
            _edit_cell(page, 1, _chips(page).index("KPI_C") + 2, "7.7")  # columns: riga, ⚠ Problemi, ...
            expect(_metric(page, "Errori")).to_have_text("0")
            # the grid reports only the edits of visible columns: hiding KPI_C and editing
            # another cell must not lose the first edit
            page.get_by_role("button", name="Remove KPI_C", exact=True).click()
            _idle(page)
            _edit_cell(page, 2, _chips(page).index("KPI_B") + 2, "99")
            expect(_metric(page, "Errori")).to_have_text("0")
            # values inferred from the file name, edited once for all the rows of the first file
            # (columns: File, Righe, Dedotti da, Valori diversi, Date, ManeuvreName, Driving_Mode, Driver);
            # the pending edit of KPI_B is committed with it
            _edit_cell(page, 1, 7, "Mario", grid=1)
            expect(_metric(page, "Errori")).to_have_text("0")

            page.locator('[data-testid="stSidebar"]').get_by_text("5 · Scrittura").click()
            _idle(page)
            _click(page, "Tutti: nuova colonna")
            _click(page, "Scrivi nel DB")
            expect(page.get_by_text("Scritti gli ID 3–10")).to_be_visible(timeout=60000)
        finally:
            browser.close()

    wb = load_workbook(demo["db"])
    rows = list(wb["KPIs"].iter_rows(values_only=True))
    records = {r[0]: dict(zip(rows[0], r)) for r in rows[1:]}
    assert records[3]["KPI_C"] == 7.7 and records[3]["IDName"] == "Dal browser"
    assert records[4]["KPI_B"] == 99
    rows = list(wb["Metadata"].iter_rows(values_only=True))
    drivers = {r[0]: dict(zip(rows[0], r))["Driver"] for r in rows[1:]}
    assert [drivers[i] for i in (3, 4, 5, 6)] == ["Mario", "Mario", "Mario", "DriverB"]
