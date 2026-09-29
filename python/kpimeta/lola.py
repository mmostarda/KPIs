"""Reading LoLa output files (port of process_excel.m).

Expected layout of each file (first sheet), as documented in process_excel.m:
  - row 1: headers (column A has no header)
  - column A: progressive number = Repetition_ID
  - completely empty rows: group separators (reset the File Name fill-down)
  - rows whose column A is empty/non-numeric: ignored
  - "File Name" can be a merged cell: its value is filled down

Differences from MATLAB, all intentional:
  - numbers stay numbers (MATLAB turned every value into text with num2str,
    keeping ~5 significant digits);
  - only columns declared `date` in the mapping are parsed as dates (MATLAB
    parsed any column whose name contained "date" or "data", emptying values
    such as "UpdateRate");
  - FileName is always kept when the File Name column exists;
  - nothing is written to disk unless export_all_db() is called.
"""

from __future__ import annotations

import datetime as dt
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import pandas as pd

from .config import Settings
from .mapping import LolaMapping
from .values import ConversionError, is_missing, looks_numeric, native, to_date, to_number, to_text
from .xlsx import open_workbook

REPETITION_COLUMN = "Repetition_ID"
FILENAME_COLUMN = "FileName"
INTERMEDIATE_FILES = {"all_db.xlsx", "metadata_all_db.xlsx"}


class LolaError(Exception):
    pass


def is_lola_output(path: str | Path) -> bool:
    """False for Excel lock files and for the intermediate files of this tool and of MATLAB
    (ALL_DB.xlsx, Metadata_ALL_DB.xlsx, *_DB.xlsx)."""
    path = Path(path)
    return (not path.name.startswith(("~$", ".")) and path.name.lower() not in INTERMEDIATE_FILES
            and not path.stem.endswith("_DB"))


@dataclass
class LolaFileReport:
    path: str
    data_rows: int = 0
    dropped_without_filename: int = 0
    columns_found: list[str] = field(default_factory=list)
    lola_names_missing: list[str] = field(default_factory=list)
    unmapped_headers: list[str] = field(default_factory=list)
    filename_header: str | None = None
    warnings: list[str] = field(default_factory=list)


def read_cells(path: str | Path) -> list[list[Any]]:
    """All values of the first sheet (like MATLAB readcell), rectangular."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        wb = open_workbook(path, LolaError, read_only=True, data_only=True)
        try:
            rows = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
        finally:
            wb.close()
    elif suffix == ".xls":
        try:
            import xlrd
        except ImportError as exc:
            raise LolaError(f"{path.name}: per i file .xls serve il pacchetto xlrd "
                            "(pip install xlrd) oppure salva il file come .xlsx") from exc
        book = xlrd.open_workbook(str(path))
        sheet = book.sheet_by_index(0)
        rows = []
        for r in range(sheet.nrows):
            row = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    row.append(None)
                else:
                    row.append(cell.value)
            rows.append(row)
    else:
        raise LolaError(f"{path.name}: formato non supportato (usa .xlsx)")
    width = max((len(r) for r in rows), default=0)
    rows = [r + [None] * (width - len(r)) for r in rows]
    # drop trailing empty columns (read-only mode can over-report the sheet size)
    while width and all(is_missing(r[width - 1]) for r in rows):
        width -= 1
        rows = [r[:width] for r in rows]
    return rows


def _repetition_id(value: Any) -> int | float | None:
    if isinstance(value, bool) or is_missing(value):
        return None
    if isinstance(value, (int, float)):
        number = float(value)
    else:
        text = to_text(value) or ""
        if not looks_numeric(text):
            return None
        number = float(text)
    return int(number) if number.is_integer() else number


def _convert(value: Any, kind: str, report: LolaFileReport, column: str, problems: dict) -> Any:
    try:
        if kind == "text":
            return to_text(value)
        if kind == "number":
            return to_number(value)
        if kind == "date":
            return to_date(value)[0]
    except ConversionError:
        problems[column] = problems.get(column, 0) + 1
        return native(value)  # kept as found: flagged later by the batch validation
    return native(value)


def parse_lola_rows(rows: list[list[Any]], mapping: LolaMapping, settings: Settings,
                    source: str = "") -> tuple[pd.DataFrame, LolaFileReport]:
    report = LolaFileReport(path=source)
    if not rows:
        report.warnings.append("file vuoto")
        return pd.DataFrame(columns=[REPETITION_COLUMN], dtype=object), report

    header = [to_text(h) or "" for h in rows[0]]
    index_of: dict[str, int] = {}
    for i, name in enumerate(header):
        key = name.strip().lower()
        if key and key not in index_of:
            index_of[key] = i

    # physical column of "File Name" (header first, then through the mapping)
    filename_idx = index_of.get(settings.lola_filename_header.strip().lower())
    if filename_idx is None:
        for entry in mapping.entries:
            if entry.db_name.lower() == FILENAME_COLUMN.lower():
                filename_idx = index_of.get(entry.lola_name.strip().lower())
                if filename_idx is not None:
                    break
    if filename_idx is None:
        report.warnings.append(f"colonna '{settings.lola_filename_header}' non trovata: FileName sarà vuoto")
    else:
        report.filename_header = header[filename_idx]

    kept: list[tuple[str, int, str]] = []
    for entry in mapping.entries:
        idx = index_of.get(entry.lola_name.strip().lower())
        if idx is None:
            report.lola_names_missing.append(entry.lola_name)
            continue
        if entry.db_name.lower() in (FILENAME_COLUMN.lower(), REPETITION_COLUMN.lower()):
            continue  # handled explicitly below
        kept.append((entry.db_name, idx, entry.type))
    used = {idx for _, idx, _ in kept} | ({filename_idx} if filename_idx is not None else set())
    report.unmapped_headers = [h for i, h in enumerate(header) if i > 0 and h.strip() and i not in used]

    data = rows[1:]
    empty_row = [all(is_missing(v) for v in r) for r in data]

    # File Name fill-down: merged cells, reset on separator rows (as in MATLAB)
    filenames: list[str | None] = []
    last: str | None = None
    for r, is_empty in zip(data, empty_row):
        raw = to_text(r[filename_idx]) if filename_idx is not None else None
        if raw and not looks_numeric(raw):
            last = raw
            filenames.append(raw)
        elif not is_empty:
            filenames.append(last)
        else:
            last = None
            filenames.append(None)

    problems: dict[str, int] = {}
    records = []
    for k, r in enumerate(data):
        rep = _repetition_id(r[0]) if r else None
        if rep is None:
            continue
        record: dict[str, Any] = {REPETITION_COLUMN: rep}
        if filename_idx is not None:
            record[FILENAME_COLUMN] = filenames[k]
        for db_name, idx, kind in kept:
            record[db_name] = _convert(r[idx], kind, report, db_name, problems)
        records.append(record)

    columns = [REPETITION_COLUMN] + ([FILENAME_COLUMN] if filename_idx is not None else [])
    columns += [db for db, _, _ in kept]
    frame = pd.DataFrame(records, columns=columns, dtype=object)
    if filename_idx is not None and len(frame):
        has_name = frame[FILENAME_COLUMN].map(lambda v: not is_missing(v))
        report.dropped_without_filename = int((~has_name).sum())
        frame = frame[has_name].reset_index(drop=True)
    report.data_rows = len(frame)
    report.columns_found = columns
    for column, count in problems.items():
        report.warnings.append(f"{column}: {count} valori non convertibili nel tipo del mapping")
    return frame, report


def read_lola_file(path: str | Path, mapping: LolaMapping, settings: Settings) -> tuple[pd.DataFrame, LolaFileReport]:
    return parse_lola_rows(read_cells(path), mapping, settings, source=str(path))


def read_lola_files(paths: list[str | Path], mapping: LolaMapping,
                    settings: Settings) -> tuple[pd.DataFrame, list[LolaFileReport]]:
    """Read several LoLa outputs and stack them (union of the columns)."""
    frames, reports = [], []
    for path in paths:
        frame, report = read_lola_file(path, mapping, settings)
        reports.append(report)
        if len(frame):
            frame = frame.copy()
            frame["_source_file"] = Path(path).name
            frames.append(frame)
    if not frames:
        return pd.DataFrame(columns=[REPETITION_COLUMN], dtype=object), reports
    stacked = pd.concat(frames, ignore_index=True, sort=False).astype(object)
    stacked = stacked.where(pd.notna(stacked), None)
    return stacked, reports


def export_all_db(frame: pd.DataFrame, path: str | Path) -> Path:
    """Write the stacked LoLa table like MATLAB ALL_DB.xlsx (Repetition_ID first)."""
    from openpyxl import Workbook

    path = Path(path)
    columns = [c for c in frame.columns if not str(c).startswith("_")]
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(columns)
    for row in frame[columns].itertuples(index=False):
        values = []
        for v in row:
            v = None if is_missing(v) else v
            if isinstance(v, dt.date) and not isinstance(v, dt.datetime):
                v = dt.datetime(v.year, v.month, v.day)
            values.append(v)
        ws.append(values)
    for cell_row in ws.iter_rows(min_row=2):
        for cell in cell_row:
            if isinstance(cell.value, dt.datetime):
                cell.number_format = "dd/mm/yyyy"
    wb.save(path)
    return path
