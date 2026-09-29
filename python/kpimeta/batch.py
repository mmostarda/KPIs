"""The batch: rows to be written, one metadata row + one KPI row each.

Sources (the "Sorgente" step of the UI):
  - LoLa outputs          -> batch_from_frame(read_lola_files(...))
  - Metadata + KPI Excel  -> batch_from_excel(...)  (e.g. Metadata_ALL_DB.xlsx)
  - manual entry          -> batch_manual(...)

Common metadata (form), templates and file-name inference are then applied by
compose(). Precedence, from strongest to weakest:
  LoLa protected columns (Repetition_ID, FileName; StartSpeed when valid)
  > file-name inference > template > form > values already in the source.
Empty values never overwrite anything.
"""

from __future__ import annotations

import copy
import datetime as dt
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import pandas as pd

from .config import Settings
from .infer import infer_from_filename, template_name
from .issues import INFO, WARNING, Issue
from .schema import Schema
from .values import ConversionError, coerce, is_missing, native, normalize_name, to_text
from .xlsx import open_workbook

LEAD_COLUMNS = ("ID", "Repetition_ID", "FileName", "StartSpeed")  # MATLAB column order
HIDDEN_PREFIX = "_"  # informative columns, never written (e.g. _source_file)
KPI_TYPE = "number"


class BatchError(Exception):
    pass


@dataclass
class Batch:
    df: pd.DataFrame
    meta_columns: list[str]
    kpi_columns: list[str]
    source: str  # lola | excel | manual
    conv_issues: dict[tuple[int, str], tuple[str, str]] = field(default_factory=dict)
    notes: list[Issue] = field(default_factory=list)
    origin: dict[tuple[int, str], str] = field(default_factory=dict)

    @property
    def n_rows(self) -> int:
        return len(self.df)

    def hidden_columns(self) -> list[str]:
        return [c for c in self.df.columns if str(c).startswith(HIDDEN_PREFIX)]

    def review_columns(self) -> list[str]:
        return [*self.meta_columns, *self.kpi_columns]

    def copy(self) -> "Batch":
        return Batch(self.df.copy(deep=True), list(self.meta_columns), list(self.kpi_columns),
                     self.source, dict(self.conv_issues), copy.deepcopy(self.notes), dict(self.origin))


def _meta_order(schema: Schema, extra: list[str]) -> list[str]:
    lead = [c for c in LEAD_COLUMNS if c in schema]
    rest = [c for c in schema.columns if c not in lead]
    return lead + rest + [c for c in extra if c not in schema]


def _empty_frame(columns: list[str], rows: int) -> pd.DataFrame:
    return pd.DataFrame({c: pd.Series([None] * rows, dtype=object) for c in columns})


def set_value(batch: Batch, row: int, column: str, raw: Any, field_type: str | None, origin: str) -> None:
    """Store a converted value; invalid values are kept aside as conversion issues."""
    try:
        value, warning = coerce(field_type, raw)
    except ConversionError as exc:
        batch.df.at[row, column] = None
        batch.conv_issues[(row, column)] = (to_text(raw) or str(raw), str(exc))
        batch.origin[(row, column)] = origin
        return
    batch.df.at[row, column] = value
    batch.conv_issues.pop((row, column), None)
    if warning:
        batch.notes.append(Issue(WARNING, warning, row, column, raw))
    if value is not None:
        batch.origin[(row, column)] = origin


def _classify(columns: list[str], schema: Schema, settings: Settings,
              meta_headers: list[str]) -> dict[str, tuple[str, str]]:
    """incoming column -> (kind, target). kind: meta | extra | kpi | skip | hidden."""
    excluded = {c.lower() for c in settings.kpi_excluded_columns}
    headers_by_norm = {normalize_name(h): h for h in meta_headers}
    out: dict[str, tuple[str, str]] = {}
    for name in columns:
        text = str(name)
        if text.startswith(HIDDEN_PREFIX):
            out[name] = ("hidden", text)
        elif (f := schema.match(text)) is not None:
            out[name] = ("meta", f.column)
        elif text.lower() in excluded:
            out[name] = ("skip", text)
        elif text in meta_headers or normalize_name(text) in headers_by_norm:
            out[name] = ("extra", text if text in meta_headers else headers_by_norm[normalize_name(text)])
        else:
            out[name] = ("kpi", text)
    return out


def batch_from_frame(frame: pd.DataFrame, schema: Schema, settings: Settings, source: str,
                     meta_headers: list[str] | None = None) -> Batch:
    """Build a batch from a table whose columns are DB names (LoLa output after
    mapping, or an imported Metadata + KPI file). Columns of the schema become
    metadata; other columns of TableMet (meta_headers) stay metadata; the rest
    are KPIs (MATLAB importMetadataKPIsFromFile)."""
    meta_headers = list(meta_headers or [])
    kinds = _classify(list(frame.columns), schema, settings, meta_headers)
    extras = list(dict.fromkeys(t for k, t in kinds.values() if k == "extra"))
    kpis = list(dict.fromkeys(t for k, t in kinds.values() if k == "kpi"))
    hidden = [t for k, t in kinds.values() if k == "hidden"]
    meta = _meta_order(schema, extras)
    rows = len(frame)
    batch = Batch(_empty_frame(meta + kpis + hidden, rows), meta, kpis, source)

    id_column = schema.id_field().column
    ignored_ids = 0
    for incoming, (kind, target) in kinds.items():
        values = list(frame[incoming])
        if kind == "meta" and target != incoming:
            batch.notes.append(Issue(INFO, f"colonna '{incoming}' letta come '{target}'", column=target))
        if kind == "skip":
            batch.notes.append(Issue(INFO, f"colonna '{incoming}' ignorata (colonna dei template)"))
            continue
        for r, raw in enumerate(values):
            if kind == "meta":
                if target == id_column:
                    ignored_ids += 0 if is_missing(raw) else 1
                    continue
                set_value(batch, r, target, raw, schema.get(target).type, source)
            elif kind == "kpi":
                set_value(batch, r, target, raw, KPI_TYPE, source)
            elif kind == "extra":
                batch.df.at[r, target] = native(raw)
                if not is_missing(raw):
                    batch.origin[(r, target)] = source
            else:
                batch.df.at[r, target] = raw
    if ignored_ids:
        batch.notes.append(Issue(INFO, "la colonna ID del file viene ignorata: "
                                       "gli ID sono assegnati al momento della scrittura", column=id_column))
    if extras:
        batch.notes.append(Issue(INFO, "colonne di TableMet non presenti nello schema, trattate come "
                                       f"metadati: {', '.join(extras)}"))
    return batch


def read_excel_table(path: str | Path) -> tuple[pd.DataFrame, int]:
    """First sheet of an Excel file, header on row 1 (like MATLAB readtable).
    Returns (frame, number of empty rows skipped)."""
    path = Path(path)
    wb = open_workbook(path, BatchError, read_only=True, data_only=True)
    try:
        rows = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    finally:
        wb.close()
    if not rows:
        raise BatchError(f"{path.name}: il file è vuoto")
    header = [to_text(h) or "" for h in rows[0]]
    keep = [i for i, h in enumerate(header) if h]
    names = [header[i] for i in keep]
    duplicates = sorted({n for n in names if names.count(n) > 1})
    if duplicates:
        raise BatchError(f"{path.name}: intestazioni ripetute: {', '.join(duplicates)}")
    data, skipped = [], 0
    for row in rows[1:]:
        row = row + [None] * (len(header) - len(row))
        values = [row[i] for i in keep]
        if all(is_missing(v) for v in values):
            skipped += 1
            continue
        data.append(values)
    return pd.DataFrame(data, columns=names, dtype=object), skipped


def batch_from_excel(path: str | Path, schema: Schema, settings: Settings,
                     meta_headers: list[str] | None = None) -> Batch:
    frame, skipped = read_excel_table(path)
    if frame.empty:
        raise BatchError(f"{Path(path).name}: nessuna riga di dati")
    batch = batch_from_frame(frame, schema, settings, "excel", meta_headers)
    if skipped:
        batch.notes.append(Issue(INFO, f"{skipped} righe vuote ignorate"))
    return batch


def batch_manual(schema: Schema, common: dict[str, Any], kpi_columns: list[str] | None = None) -> Batch:
    meta = _meta_order(schema, [])
    kpis = list(dict.fromkeys(kpi_columns or []))
    batch = Batch(_empty_frame(meta + kpis, 1), meta, kpis, "manual")
    apply_common(batch, schema, common, mode="overwrite")
    return batch


def apply_common(batch: Batch, schema: Schema, values: dict[str, Any], mode: str = "overwrite") -> None:
    """Form values on every row. mode='fill' only fills empty cells."""
    protected = batch.source != "manual"
    for column, raw in values.items():
        f = schema.get(column)
        if f is None or f.type == "id" or column not in batch.df.columns:
            continue
        try:
            value, _ = coerce(f.type, raw)
        except ConversionError:
            continue  # the form validates its own values
        if value is None:
            continue
        for r in range(batch.n_rows):
            current = batch.df.at[r, column]
            if protected and f.lola_policy == "keep":
                continue
            if protected and f.lola_policy == "keep_if_valid" and current is not None:
                continue
            if mode == "fill" and (current is not None or (r, column) in batch.conv_issues):
                continue
            set_value(batch, r, column, value, f.type, "form")


def apply_templates(batch: Batch, schema: Schema, settings: Settings, templates: pd.DataFrame) -> int:
    """Row by row: template named after the first N tokens of FileName
    (IDName_Template). Returns the number of rows that found a template."""
    name_column = next((c for c in templates.columns
                        if str(c).strip().lower() == settings.template_name_column.lower()), None)
    if name_column is None or "FileName" not in batch.df.columns:
        batch.notes.append(Issue(WARNING, f"template non applicabili: colonna "
                                          f"'{settings.template_name_column}' o FileName mancante"))
        return 0
    skip = {str(name_column).lower(), settings.template_id_column.lower()}
    index: dict[str, int] = {}
    for i, name in enumerate(templates[name_column]):
        text = to_text(name)
        if text:
            index.setdefault(text.lower(), i)
    missing: dict[str, list[int]] = {}
    applied = 0
    for r in range(batch.n_rows):
        name = template_name(batch.df.at[r, "FileName"], settings)
        if not name:
            continue
        i = index.get(name.lower())
        if i is None:
            missing.setdefault(name, []).append(r)
            continue
        applied += 1
        for tcol in templates.columns:
            if str(tcol).lower() in skip:
                continue
            raw = templates.iloc[i][tcol]
            if is_missing(raw):
                continue
            f = schema.match(str(tcol))
            if f is None:
                if tcol not in batch.df.columns:
                    batch.df[tcol] = pd.Series([None] * batch.n_rows, dtype=object)
                    batch.meta_columns.append(tcol)
                batch.df.at[r, tcol] = native(raw)
                batch.origin[(r, tcol)] = "template"
                continue
            if f.type == "id" or f.lola_policy == "keep":
                continue
            current_origin = batch.origin.get((r, f.column))
            if (f.lola_policy == "keep_if_valid" and batch.df.at[r, f.column] is not None
                    and current_origin in ("lola", "excel")):
                continue
            set_value(batch, r, f.column, raw, f.type, "template")
    for name, rows in missing.items():
        batch.notes.append(Issue(WARNING, f"template '{name}' non trovato in {settings.templates_sheet}: "
                                          f"{len(rows)} righe restano con i valori del form",
                                 row=rows[0], column="FileName"))
    return applied


def apply_filename_inference(batch: Batch, schema: Schema, settings: Settings) -> int:
    """Date, Driving_Mode, ManeuvreName, Driver from FileName. Returns rows inferred."""
    if "FileName" not in batch.df.columns:
        return 0
    inferred = 0
    for r in range(batch.n_rows):
        values, problem = infer_from_filename(batch.df.at[r, "FileName"], settings, schema)
        if problem:
            batch.notes.append(Issue(WARNING, f"nome file non riconosciuto: {problem}", r, "FileName"))
        if values:
            inferred += 1
        for column, value in values.items():
            f = schema.get(column)
            if f is not None and column in batch.df.columns:
                set_value(batch, r, column, value, f.type, "filename")
    return inferred


def compose(base: Batch, schema: Schema, settings: Settings, common: dict[str, Any], *,
            mode: str = "overwrite", templates: pd.DataFrame | None = None,
            infer_template: bool = False, infer_filename: bool = False) -> Batch:
    """New batch = base + form values + templates + file-name inference."""
    batch = base.copy()
    apply_common(batch, schema, common, mode=mode)
    if infer_template and templates is not None:
        apply_templates(batch, schema, settings, templates)
    if infer_filename:
        apply_filename_inference(batch, schema, settings)
    return batch


def assign_ids(batch: Batch, start_id: int, id_column: str = "ID") -> None:
    if id_column in batch.df.columns:
        batch.df[id_column] = pd.Series(range(start_id, start_id + batch.n_rows), dtype=object)


def export_batch_excel(batch: Batch, path, id_column: str = "ID"):
    """Write the batch like Metadata_ALL_DB.xlsx (re-importable with batch_from_excel).
    `path` can also be a binary file-like object."""
    from openpyxl import Workbook

    path = Path(path) if isinstance(path, str) else path
    columns = [c for c in batch.review_columns() if c != id_column]
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(columns)
    for r in range(batch.n_rows):
        row = []
        for c in columns:
            v = batch.df.at[r, c]
            if (r, c) in batch.conv_issues and v is None:
                v = batch.conv_issues[(r, c)][0]  # keep the original text so it can be fixed
            if isinstance(v, dt.date) and not isinstance(v, dt.datetime):
                v = dt.datetime(v.year, v.month, v.day)
            row.append(v)
        ws.append(row)
    for cells in ws.iter_rows(min_row=2):
        for cell in cells:
            if isinstance(cell.value, dt.datetime):
                cell.number_format = "dd/mm/yyyy"
    wb.save(path)
    return path
