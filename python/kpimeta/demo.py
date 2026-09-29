"""SYNTHETIC demo data (DB, LoLa outputs, mapping) for tests and for trying the app.

The layout follows what the MATLAB code expects; the values are invented and
must not be taken as real: validate the tool on a copy of the real DB and on
real LoLa outputs before using it.
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.cell import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .config import load_settings
from .schema import Schema, load_schema

SUPPORT_VALUES = {
    ("Cluster", "Cluster"): ["Cluster A", "Cluster B"],
    ("SubCluster", "SubCluster"): ["Sub 1", "Sub 2"],
    ("Brand", "Brand"): ["Brand X", "Brand Y"],
    ("Model", "Model"): ["Model 1", "Model 2"],
    ("Maneuvre", "Maneuvre"): ["BrakeTest", "Slalom"],
    ("Condition", "Condition"): ["dry", "wet"],
    ("Surface", "Surface"): ["asphalt", "ice"],
    ("Tire", "Tire"): ["Summer", "Winter"],
    ("Brake", "BrakeType"): ["Type 1", "Type 2"],
    ("SmT_SW_TireID", "TireID"): ["T-01", "T-02"],
    ("SmT_SW_Estimate", "Condition"): ["ON", "OFF"],
    ("Susp", "SuspActiveStatus"): ["Active", "Passive"],
    ("Steer", "SteeringType"): ["EPS", "SbW"],
    ("BL_Phase", "Phase"): ["Phase 1", "Phase 2"],
}

KPI_COLUMNS = ["KPI_A", "KPI_B", "KPI_C"]
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="4472C4")


def _add_table(ws, name: str, columns: list[str], rows: list[list], style: str = "TableStyleMedium2") -> Table:
    ws.append(columns)
    for row in rows:
        ws.append(row)
    last_row = max(2, len(rows) + 1)  # an Excel table has at least one data row
    ref = f"A1:{get_column_letter(len(columns))}{last_row}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True)
    ws.add_table(table)
    for i in range(1, len(columns) + 1):
        ws.column_dimensions[get_column_letter(i)].width = max(10, min(28, len(columns[i - 1]) + 2))
    return table


def make_demo_db(path: str | Path, schema: Schema, existing_rows: int = 2) -> Path:
    path = Path(path)
    wb = Workbook()
    ws = wb.active
    ws.title = "Metadata"
    samples = {
        "IDName": "Demo {i}", "Date": "0{i}/02/2026", "Cluster": "Cluster A", "Sub_Cluster": "Sub 1",
        "Brand": "Brand X", "Model": "Model 1", "ManeuvreName": "BrakeTest", "StartSpeed": 100,
        "Repetition_ID": "{i}", "Track": "Track 1", "Surface": "asphalt", "Condition": "dry",
        "Driving_Mode": "Sport", "ABS": "ON", "ESC": "ON", "TCS": "OFF", "Driver": "DriverA",
        "FileName": "DEMO_OLD_2026020{i}_R0{i}_Sport_BrakeTest_DriverA_00{i}.mf4",
    }
    rows = []
    for i in range(1, existing_rows + 1):
        row = []
        for column in schema.columns:
            value = samples.get(column)
            if column == schema.id_field().column:
                value = i
            elif isinstance(value, str):
                value = value.format(i=i)
                value = int(value) if column == "Repetition_ID" else value
            row.append(value)
        rows.append(row)
    _add_table(ws, "TableMet", schema.columns, rows)

    kpi = wb.create_sheet("KPIs")
    kpi_rows = [[i, f"Demo {i}", f"0{i}/02/2026", 1.1 * i, 2.2 * i, None] for i in range(1, existing_rows + 1)]
    _add_table(kpi, "TableKPI", ["ID", "IDName", "Date", *KPI_COLUMNS], kpi_rows)

    templates = wb.create_sheet("Metadata_Templates")
    columns = ["ID_Template", "IDName_Template", "Cluster", "Sub_Cluster", "Brand", "Model", "Track",
               "Surface", "Condition", "ABS", "ESC", "TCS", "T_Front_type", "T_Rear_type", "B_F_Cal_Date"]
    templates.append(columns)
    templates.append([1, "PRJ1_VEH1", "Cluster A", "Sub 1", "Brand X", "Model 1", "Track 1", "asphalt",
                      "dry", "ON", "ON", "ON", "Summer", "Summer", dt.datetime(2026, 1, 10)])
    templates.append([2, "PRJ2_VEH2", "Cluster B", "Sub 2", "Brand Y", "Model 2", "Track 2", "ice",
                      "wet", "ON", "OFF", "OFF", "Winter", "Winter", None])
    templates["O2"].number_format = "dd/mm/yyyy"
    for cell in templates[1]:
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL

    for (sheet, column), values in SUPPORT_VALUES.items():
        if sheet in wb.sheetnames:
            continue
        ws_support = wb.create_sheet(sheet)
        _add_table(ws_support, f"T_{sheet}", ["ID", column], [[i + 1, v] for i, v in enumerate(values)])
    wb.save(path)
    return path


def _lola_file(path: Path, groups: list[tuple[str, list[list]]], extra_headers: list[str]) -> Path:
    """One LoLa output: column A = repetition, File Name merged over each group,
    a statistics row after each group and an empty separator row."""
    wb = Workbook()
    ws = wb.active
    ws.title = "LoLa"
    ws.append([None, "File Name", "Speed start [km/h]", "Kpi A", "Kpi B", "Kpi C", *extra_headers])
    row = 2
    for filename, values in groups:
        first = row
        for rep, data in enumerate(values, start=1):
            ws.append([rep, filename if rep == 1 else None, *data])
            row += 1
        if len(values) > 1:
            ws.merge_cells(start_row=first, start_column=2, end_row=row - 1, end_column=2)
        ws.append(["Mean", None, None, None, None, None])  # statistics row: ignored
        ws.append([])  # separator
        row += 2
    wb.save(path)
    return path


def make_demo_lola(folder: str | Path) -> list[Path]:
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    first = _lola_file(folder / "LoLa_output_1.xlsx", [
        ("PRJ1_VEH1_20260315_R01_Sport_BrakeTest_DriverA_001.mf4",
         [[100.2, 0.123456789, 12.5, "n.a.", 1], [99.8, 0.223456789, 13.25, 7.5, 2], [None, 0.3, 14.0, "NA", 3]]),
        ("PRJ1_VEH1_20260316_R02_Comfort_Slalom_DriverB_002.mf4",
         [[80.0, 1.5, 2.5, 3.5, 4], [80.5, 1.6, 2.6, 3.6, 5]]),
    ], extra_headers=["Kpi D"])
    second = _lola_file(folder / "LoLa_output_2.xlsx", [
        ("PRJ2_VEH2_20260320_R01_Eco_BrakeTest_DriverC_003.mf4", [[60.0, 5.5, 6.5, 7.5, 8]]),
        ("PRJ9_VEH9_20260321_R01_Eco_BrakeTest_DriverC_004.mf4", [[61.0, 5.6, 6.6, 7.6, 9]]),
        ("shortname.mf4", [[62.0, 5.7, 6.7, 7.7, 10]]),
    ], extra_headers=["Kpi D"])
    return [first, second]


def make_demo_mapping(path: str | Path) -> Path:
    path = Path(path)
    path.write_text(
        "KPI_LoLa_Name;KPI_DB_Name;Type\n"
        "File Name;FileName;text\n"
        "Speed start [km/h];StartSpeed;number\n"
        "Kpi A;KPI_A;\n"
        "Kpi B;KPI_B;\n"
        "Kpi C;KPI_C;\n"
        "Kpi D;KPI_D;\n",
        encoding="utf-8")
    return path


def make_demo(folder: str | Path, config_dir: str | Path | None = None) -> dict[str, Path | list[Path]]:
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    settings = load_settings(config_dir)
    schema = load_schema(settings.schema_path)
    return {
        "db": make_demo_db(folder / "KPI_DB_demo.xlsx", schema),
        "lola": make_demo_lola(folder / "lola"),
        "mapping": make_demo_mapping(folder / "lola_mapping_demo.csv"),
    }
