"""Opening Excel files: a damaged or non-Excel file gives a readable error, not a traceback."""

from __future__ import annotations

import zipfile
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


def invalid_file_message(path: str | Path) -> str:
    return f"{Path(path).name}: non è un file Excel valido o è danneggiato"


def open_workbook(path: str | Path, error: type[Exception], **options: Any):
    """load_workbook(path, **options), raising `error` when the file is not a valid workbook."""
    try:
        return load_workbook(path, **options)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, EOFError) as exc:
        raise error(invalid_file_message(path)) from exc
